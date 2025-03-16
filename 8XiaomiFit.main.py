# 8XiaomiFit.main.py 创建每天日期的各种健康数据

import re  # 用于正则表达式，匹配日志中的特定数据模式
import json  # 用于解析日志中嵌入的 JSON 格式数据（如睡眠数据）
import mysql.connector  # 用于连接和操作 MySQL 数据库
import openpyxl  # 用于创建和操作 Excel 文件（直接操作工作簿）
import pandas as pd  # 提供数据处理功能，虽然这里未直接使用，但可能是备用工具
from datetime import datetime  # 用于处理日期和时间，生成时间戳或格式化日期
from concurrent.futures import ThreadPoolExecutor, as_completed  # 用于多线程并行处理，提高解析效率
import logging  # 用于记录程序运行日志，便于调试和跟踪

# 设置日志
logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(levelname)s - %(message)s',
    handlers=[  # 设置日志输出格式
        logging.FileHandler("program.log", encoding='utf-8'),
        # logging.StreamHandler()  # 输出到控制台
    ]
)

# 文件路径配置
CURRENT_DATE = datetime.now().strftime('%Y_%m_%d')  # 当前日期格式为：2025_02_22，便于创建文件夹和文件
LOG_FILE_PATH = r'C:\Users\maikojoke\Desktop\小米手环数据自动化导出代码\health\apps\com.mi.health\ef\log\XiaomiFit.main.log'  # 日志文件路径
EXCEL_OUTPUT_PATH = f'C:\\Users\\maikojoke\\Desktop\\小米手环数据自动化导出代码\\craw file\\{CURRENT_DATE}_XiaomiFit.main.xlsx'  # Excel 输出文件路径

# MySQL 配置
DB_CONFIG = {
    "host": "localhost",  # MySQL 数据库主机地址
    "user": "root",  # MySQL 用户名
    "password": "123456",  # MySQL 密码
    "database": "daily_data"  # MySQL 数据库名
}

# 正则表达式模式
PATTERNS = {
    "daily_goal": re.compile(r"(\d{4}-\d{2}-\d{2} \d{2}:\d{2}:\d{2})\.\d+\|.*?RainbowReport, mDailyGoalReport = DailyGoalReport\(time=\d+, time=[\d\-]+\s[\d:]+, goals=\[(.*?)\]\)"),  # 目标数据正则模式
    "sleep": re.compile(r"(\d{4}-\d{2}-\d{2} \d{2}:\d{2}:\d{2})\.\d+\|.*?value='(.*?)'"),  # 睡眠数据正则模式
    "heart_rate": re.compile(r"HrItem\(sid=\d+, time=(\d+), hr=(\d+)\)"),  # 心率数据正则模式
    "step": re.compile(r"StepItem\(sid=\d+, time=(\d+), steps=(\d+), distance=(\d+), calories=(\d+\.\d+)"),  # 步数数据正则模式
    "calorie": re.compile(r"CaloriesItem\(time=(\d+), calories = (\d+\.\d+)"),  # 卡路里数据正则模式
    "stress": re.compile(r"StressItem\(sid=\d+, time=(\d+), stress=(\d+)"),  # 压力数据正则模式
    "spo2": re.compile(r"Spo2Item\(time=(\d+), sid=\d+, spo2=(\d+)"),  # 血氧数据正则模式
    "vitality": re.compile(r"VitalityItem\(time=(\d+), lowVitality=(\d+), mediumVitality=(\d+), highVitality=(\d+), totalVitality=(\d+), activityType=(\d+), activityDuration=(\d+)"),  # 活跃度数据正则模式
    "weight": re.compile(r"WeightItem\(time=(\d+),.*?weight=(\d+\.\d+)\)"),  # 体重数据正则模式
    "valid_stand": re.compile(r"StandRecord\(time=(\d+), standCount=(\d+)\)"),  # 站立数据正则模式
}

# 数据类型和字段定义
DATA_TYPES = {
    "daily_goal": ["timestamp", "goal_type_name", "goal_value", "achieved_value"],  # 每日目标数据字段
    "sleep": ["timestamp", "avg_hr", "avg_spo2", "sleep_awake_duration", "sleep_deep_duration", "sleep_light_duration", "sleep_rem_duration", "sleep_score", "max_hr", "max_spo2", "min_hr", "min_spo2", "duration"],  # 睡眠数据字段
    "heart_rate": ["timestamp", "heart_rate"],  # 心率数据字段
    "step": ["timestamp", "steps", "distance_m", "calories"],  # 步数数据字段
    "calorie": ["timestamp", "calories"],  # 卡路里数据字段
    "stress": ["timestamp", "stress_level"],  # 压力数据字段
    "spo2": ["timestamp", "spo2"],  # 血氧数据字段
    "vitality": ["timestamp", "low_vitality", "medium_vitality", "high_vitality", "total_vitality", "activity_type", "activity_duration"],  # 活跃度数据字段
    "weight": ["timestamp", "weight_kg"],  # 体重数据字段
    "valid_stand": ["timestamp", "stand_count"],  # 站立数据字段
}

# 创建 Excel 工作簿
def initialize_excel():
    wb = openpyxl.Workbook()  # 创建新的 Excel 工作簿
    wb.remove(wb.active)  # 移除默认的空工作表
    for data_type in DATA_TYPES:  # 遍历所有数据类型
        ws = wb.create_sheet(title=f"{CURRENT_DATE}_{data_type}")  # 使用带下划线的日期格式创建工作表
        ws.append([f"{data_type}_{field}" for field in DATA_TYPES[data_type]])  # 向工作表中添加字段名作为表头
    return wb  # 返回创建好的工作簿

# 数据提取函数
def extract_data(entry, data_type):
    try:
        if data_type == "daily_goal":
            timestamp, goals = entry.groups()  # 提取目标数据中的时间戳和目标数据
            goal_items = re.findall(r"GoalItem\(type=\d+,\s*type=(\w+),\s*goal=(\d+),\s*achieved=(\d+)\)", goals)  # 提取每个目标项的详细信息
            return [(timestamp, item[0], item[1], item[2]) for item in goal_items] if goal_items else []  # 返回目标项的列表
        elif data_type == "sleep":
            timestamp, sleep_data = entry.groups()  # 提取睡眠数据中的时间戳和嵌套的 JSON 数据
            if not sleep_data:  # 检查空值
                logging.warning(f"睡眠数据为空: {timestamp}")
                return [timestamp] + [None] * (len(DATA_TYPES["sleep"]) - 1)
            try:
                sleep_data_dict = json.loads(sleep_data.replace("'", "\""))  # 解析 JSON 数据
                if isinstance(sleep_data_dict, dict):  # 确保是字典
                    return [
                        timestamp, sleep_data_dict.get("avg_hr", None), sleep_data_dict.get("avg_spo2", None),
                        sleep_data_dict.get("sleep_awake_duration", None),
                        sleep_data_dict.get("sleep_deep_duration", None),
                        sleep_data_dict.get("sleep_light_duration", None),
                        sleep_data_dict.get("sleep_rem_duration", None),
                        sleep_data_dict.get("sleep_score", None), sleep_data_dict.get("max_hr", None),
                        sleep_data_dict.get("max_spo2", None), sleep_data_dict.get("min_hr", None),
                        sleep_data_dict.get("min_spo2", None), sleep_data_dict.get("total_duration", None)
                    ]
                else:
                    logging.error(f"解析后的睡眠数据不是字典类型: {sleep_data_dict}")
                    return [timestamp] + [None] * (len(DATA_TYPES["sleep"]) - 1)
            except json.JSONDecodeError as e:
                logging.error(f"JSON解析错误 ({data_type}): {e}, 原数据: {sleep_data}")
                return [timestamp] + [None] * (len(DATA_TYPES["sleep"]) - 1)
        else:
            groups = list(entry.groups())  # 提取其他类型的日志数据
            timestamp = datetime.fromtimestamp(int(groups[0])).strftime('%Y-%m-%d %H:%M:%S')  # 将时间戳转换为可读的时间格式
            return [timestamp] + groups[1:]  # 返回时间戳和其他数据
    except (json.JSONDecodeError, ValueError) as e:
        logging.error(f"数据提取错误 ({data_type}): {e}")  # 如果解析数据出错，记录错误日志
        return [None] * len(DATA_TYPES[data_type])  # 返回空数据

# 多线程解析日志（添加时间戳去重）
def parse_log_file(log_data, data_type):
    pattern = PATTERNS[data_type]  # 获取当前数据类型的正则模式
    entries = pattern.finditer(log_data)  # 查找所有符合模式的日志条目
    data_list = []
    seen_timestamps = set()  # 用于记录已见过的时间戳
    for entry in entries:  # 遍历每个日志条目
        logging.info(f"匹配的 {data_type} 数据: {entry.group()}")  # 记录匹配的日志
        data = extract_data(entry, data_type)  # 提取数据
        if data_type == "daily_goal":
            for item in data:
                if item[0] and datetime.strptime(item[0], '%Y-%m-%d %H:%M:%S').strftime('%Y%m%d') == CURRENT_DATE.replace('_', ''):
                    if item[0] not in seen_timestamps:  # 检查时间戳是否重复
                        seen_timestamps.add(item[0])
                        data_list.append(item)
        else:
            if data[0] and datetime.strptime(data[0], '%Y-%m-%d %H:%M:%S').strftime('%Y%m%d') == CURRENT_DATE.replace('_', ''):
                if data[0] not in seen_timestamps:  # 检查时间戳是否重复
                    seen_timestamps.add(data[0])
                    data_list.append(data)

    # 按时间戳排序
    data_list.sort(key=lambda x: x[0])  # 排序数据列表
    return data_type, data_list  # 返回数据类型和对应的数据

# 创建数据库表
def create_table(cursor, data_type):
    table_name = f"{CURRENT_DATE}_{data_type}"  # 使用带下划线的日期格式创建表名
    fields = [f"{data_type}_{field}" for field in DATA_TYPES[data_type]]  # 将数据字段名附加到表格定义中
    create_table_sql = f"""
    CREATE TABLE IF NOT EXISTS `{table_name}` (  # 创建 SQL 语句
        id INT AUTO_INCREMENT PRIMARY KEY,
        {fields[0]} DATETIME NOT NULL UNIQUE,  # 时间戳字段
        {', '.join([f"{field} VARCHAR(255)" for field in fields[1:]])}  # 其他字段类型
    ) ENGINE=InnoDB DEFAULT CHARSET=utf8;
    """
    cursor.execute(create_table_sql)  # 执行创建表的 SQL 语句
    logging.info(f"数据库表 `{table_name}` 已创建或已存在。")  # 记录日志

# 插入数据到数据库
def insert_data_to_db(cursor, data_type, data_list):
    if not data_list:  # 如果数据列表为空，直接返回
        return
    table_name = f"{CURRENT_DATE}_{data_type}"  # 使用带下划线的日期格式创建表名
    columns = [f"{data_type}_{field}" for field in DATA_TYPES[data_type]]  # 获取表格的列名
    placeholders = ', '.join(['%s'] * len(columns))  # 为 SQL 插入语句准备占位符
    update_clause = ', '.join([f"{col} = VALUES({col})" for col in columns[1:]])  # 为 ON DUPLICATE KEY UPDATE 子句创建更新语句
    insert_sql = f"""
    INSERT INTO `{table_name}` ({', '.join(columns)})
    VALUES ({placeholders})
    ON DUPLICATE KEY UPDATE {update_clause}
    """
    for row in data_list:
        cursor.execute(insert_sql, tuple(row))  # 执行插入语句

# 确保秒数唯一性，避免重复
def is_duplicate_entry(cursor, data_type, timestamp):
    table_name = f"{CURRENT_DATE}_{data_type}"  # 使用带下划线的日期格式创建表名
    check_sql = f"SELECT COUNT(*) FROM `{table_name}` WHERE `{data_type}_timestamp` = %s"  # 检查是否已存在相同的时间戳
    cursor.execute(check_sql, (timestamp,))
    count = cursor.fetchone()[0]  # 获取匹配的记录数
    return count > 0  # 如果已经存在，返回 True

# 主函数（确保一致性）
def main():
    # 读取日志文件
    try:
        with open(LOG_FILE_PATH, 'r', encoding='utf-8') as log_file:
            log_data = log_file.read()  # 读取日志文件内容
    except Exception as e:
        logging.error(f"读取日志文件失败: {e}")  # 如果读取失败，记录错误日志
        return

    # 初始化 Excel 文件
    wb = initialize_excel()  # 初始化工作簿
    cursor = None
    connection = None
    try:
        # 连接数据库
        connection = mysql.connector.connect(**DB_CONFIG)  # 创建数据库连接
        cursor = connection.cursor()  # 创建游标

        # 创建表格
        for data_type in DATA_TYPES:
            create_table(cursor, data_type)  # 创建每个数据类型的数据库表

        # 使用线程池并行解析数据
        with ThreadPoolExecutor() as executor:
            futures = {executor.submit(parse_log_file, log_data, data_type): data_type for data_type in DATA_TYPES}  # 使用多线程解析日志文件
            for future in as_completed(futures):  # 等待线程执行完毕
                data_type, data_list = future.result()  # 获取解析结果
                if data_list:
                    # 先插入数据库
                    insert_data_to_db(cursor, data_type, data_list)  # 将数据插入数据库
                    # 再写入 Excel，确保一致性
                    sheet = wb[f"{CURRENT_DATE}_{data_type}"]  # 获取当前数据类型的工作表
                    for row in data_list:
                        sheet.append(row)  # 将数据行添加到工作表中

        # 保存 Excel 文件
        wb.save(EXCEL_OUTPUT_PATH)  # 保存工作簿到指定路径
        logging.info(f"数据已成功导出到 Excel：{EXCEL_OUTPUT_PATH}")  # 记录成功日志

        # 提交事务
        connection.commit()  # 提交数据库事务
    except mysql.connector.Error as err:
        logging.error(f"MySQL 错误: {err}")  # 记录 MySQL 错误
    except Exception as e:
        logging.error(f"发生错误: {e}")  # 记录其他错误
    finally:
        if cursor:
            cursor.close()  # 关闭游标
        if connection:
            connection.close()  # 关闭数据库连接
            print("数据库连接已关闭。")
if __name__ == '__main__':
    main()  # 启动主程序
